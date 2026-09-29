"""Pull text out of uploaded documents, one entry per page.

Supported:
    .pdf                    text layer via pdfium; scanned pages (no text layer) are OCR'd
    .docx                   paragraphs and tables
    .txt / .md              read as-is
    .png .jpg .jpeg .webp .tif .tiff    OCR (photos or scans of warranty cards, bills...)

OCR (backend/parser/ocr.py) uses Tesseract, a separate program: `brew install tesseract`
on macOS. Without it, text-based PDFs, Word and text files still work; images and scanned
pages are reported as needing OCR. OCR'd pages keep their word positions and a confidence
score (Page.ocr_data) so answers can be highlighted on scans and blurry photos flagged.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

import ftfy

from backend.parser.ocr import OcrPage, OcrUnavailableError, ocr_available, read_image, require_ocr

TEXT_SUFFIXES = {".txt", ".md"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".heic", ".heif"}  # HEIC: iPhones, many Androids
SUPPORTED_SUFFIXES = {".pdf", ".docx"} | TEXT_SUFFIXES | IMAGE_SUFFIXES

MIN_PAGE_CHARS = 25  # a PDF page with less text than this is treated as a scan
OCR_RENDER_SCALE = 3  # render scanned PDF pages at 216 dpi before OCR


class UnsupportedFileError(ValueError):
    pass


_HEIF_BRANDS = {b"heic", b"heix", b"hevc", b"hevx", b"heim", b"heis", b"mif1", b"msf1"}


def sniff(path: str | Path) -> str | None:
    """The file's real type from its first bytes, as a suffix (".pdf", ".jpg"…), or None.

    Phones and chat apps often send files with no extension ("Aadhaar"), a wrong one (a JPEG called
    ".pdf") or an unusual one (".JPG", ".jfif"); the contents say what it is."""
    with open(path, "rb") as f:
        head = f.read(2048)
    if b"%PDF-" in head[:1024]:
        return ".pdf"
    if head[:3] == b"\xff\xd8\xff":
        return ".jpg"
    if head[:8] == b"\x89PNG\r\n\x1a\n":
        return ".png"
    if head[:4] == b"RIFF" and head[8:12] == b"WEBP":
        return ".webp"
    if head[:4] in (b"II*\x00", b"MM\x00*"):
        return ".tif"
    if head[4:8] == b"ftyp" and head[8:12] in _HEIF_BRANDS:
        return ".heic"
    if head[:4] == b"PK\x03\x04":
        import zipfile

        try:
            with zipfile.ZipFile(path) as z:
                return ".docx" if "word/document.xml" in z.namelist() else None
        except zipfile.BadZipFile:
            return None
    if head[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        raise UnsupportedFileError("This is an old Word (.doc) or Excel file. Open it and save it as PDF or .docx, "
                                   "then upload that.")
    if head:
        try:
            sample = head.decode("utf-8")
        except UnicodeDecodeError:
            try:  # a multi-byte character cut at the end of the sample
                sample = head[:-3].decode("utf-8")
            except UnicodeDecodeError:
                return None
        if sum(ch.isprintable() or ch in "\n\r\t" for ch in sample) / len(sample) > 0.95:
            return ".txt"
    return None


def file_type(path: str | Path, name: str = "") -> str:
    """The suffix to read a file as: its contents first, then its name. Raises UnsupportedFileError."""
    found = sniff(path)
    suffix = Path(name or str(path)).suffix.lower()
    if found == ".txt" and suffix and suffix not in TEXT_SUFFIXES:
        found = None  # text-looking bytes in a ".xlsx" / ".csv"…: not a document we read
    if found:
        return found
    if suffix in SUPPORTED_SUFFIXES and suffix not in (".pdf", ".docx", *IMAGE_SUFFIXES):
        return suffix  # .md: the contents looked binary, but trust the name for text
    raise UnsupportedFileError(
        "This file isn't a PDF, photo, Word (.docx) or text file, so it can't be read. "
        "If it's a document, save or print it as PDF and upload that.")


class PasswordRequiredError(ValueError):
    """A locked PDF (e-Aadhaar, e-PAN, bank statements): ask for its password."""

    def __init__(self, wrong: bool):
        self.wrong = wrong
        super().__init__("That password didn't open this PDF. Try again." if wrong else
                         "This PDF is locked with a password. Enter it to add the document (e-Aadhaar: the first "
                         "4 letters of your name in capitals + year of birth; e-PAN: date of birth as DDMMYYYY).")


def open_pdf(path, password: str | None = None):
    import pypdfium2 as pdfium

    try:
        return pdfium.PdfDocument(path, password=password or None)
    except pdfium.PdfiumError as e:
        if "password" in str(e).lower():
            raise PasswordRequiredError(wrong=bool(password)) from e
        raise


def unlock_pdf(path: Path, password: str, dest: Path) -> Path:
    """Save an unlocked copy, so the vault can show and search it later without the password.
    (The copy stays inside the vault, which is where the protection lives.)"""
    import pypdfium2.raw as raw

    pdf = open_pdf(path, password)
    try:
        pdf.save(str(dest), flags=raw.FPDF_REMOVE_SECURITY)
    finally:
        pdf.close()
    return dest


MIN_IMAGE_SHARE = 0.4  # a page this much covered by pictures is also read with OCR


def _image_share(page) -> float:
    """How much of the page is covered by pictures (a photo of an ID card pasted into a PDF)."""
    import pypdfium2.raw as raw

    width, height = page.get_size()
    covered = 0.0
    try:
        for obj in page.get_objects(filter=[raw.FPDF_PAGEOBJ_IMAGE], max_depth=2):
            left, bottom, right, top = obj.get_bounds()
            covered += max(0.0, right - left) * max(0.0, top - bottom)
    except Exception:
        return 0.0
    return min(1.0, covered / (width * height)) if width and height else 0.0


@dataclass
class Page:
    number: int  # 1-based
    text: str
    ocr: bool = False
    ocr_data: OcrPage | None = None  # word boxes and confidence, for OCR'd pages


def clean_text(text: str) -> str:
    """Repair mis-encoded characters and normalize whitespace, keeping line breaks."""
    text = ftfy.fix_text(text.replace("\x00", ""))
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _pdf_pages(path: Path, password: str | None = None) -> list[Page]:
    pages = []
    pdf = open_pdf(path, password)
    try:
        for i, page in enumerate(pdf, start=1):
            textpage = page.get_textpage()
            text = clean_text(textpage.get_text_range())
            textpage.close()
            data = None
            scanned = len(text) < MIN_PAGE_CHARS
            # A page with a little text around a big picture (an ID card pasted into a PDF, a scan with a
            # typed header) is read with OCR too, and keeps whichever reading has more to say.
            mostly_picture = not scanned and len(text) < 600 and _image_share(page) >= MIN_IMAGE_SHARE
            if (scanned or mostly_picture) and ocr_available():
                read = read_image(page.render(scale=OCR_RENDER_SCALE).to_pil())
                if scanned or len(clean_text(read.text)) > len(text):
                    data, text = read, clean_text(read.text)
            page.close()
            pages.append(Page(i, text, data is not None, data))
    finally:
        pdf.close()
    if not any(p.text for p in pages) and not ocr_available():
        raise OcrUnavailableError(
            "This PDF has no text layer (it's a scan). Install Tesseract (macOS: brew install tesseract) "
            "and restart the app."
        )
    return pages


def _docx_pages(path: Path) -> list[Page]:
    import docx

    doc = docx.Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return [Page(1, clean_text("\n".join(parts)))]


def _image_pages(path: Path) -> list[Page]:
    from PIL import Image

    require_ocr()  # say what to install before failing on anything else
    try:  # HEIC photos (iPhones, many Androids), whatever the file is called
        import pillow_heif

        pillow_heif.register_heif_opener()
    except ImportError:
        pass
    with Image.open(path) as image:
        data = read_image(image)
    return [Page(1, clean_text(data.text), ocr=True, ocr_data=data)]


def extract_pages(path: str | Path, password: str | None = None) -> list[Page]:
    """Return the text of each page. Raises UnsupportedFileError / OcrUnavailableError /
    PasswordRequiredError (a locked PDF: pass its `password`)."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES or sniff(path) not in (None, suffix, ".txt"):
        suffix = file_type(path)  # no extension, or the wrong one: go by the contents
    if suffix == ".pdf":
        return _pdf_pages(path, password)
    if suffix == ".docx":
        return _docx_pages(path)
    if suffix in TEXT_SUFFIXES:
        return [Page(1, clean_text(path.read_text(encoding="utf-8", errors="replace")))]
    if suffix in IMAGE_SUFFIXES:
        return _image_pages(path)
    raise UnsupportedFileError(
        f"Unsupported file type {suffix or '(none)'}. Supported: {', '.join(sorted(SUPPORTED_SUFFIXES))}"
    )
