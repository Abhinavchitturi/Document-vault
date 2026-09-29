import docx
import pytest
from PIL import Image

from backend.parser import extract, ocr
from backend.parser.categories import guess_category
from backend.parser.extract import OcrUnavailableError, UnsupportedFileError, clean_text, extract_pages
from tests.pdf_helper import make_pdf


def test_pdf_pages_are_extracted_separately(tmp_path):
    path = tmp_path / "policy.pdf"
    path.write_bytes(make_pdf(["Policy expires on 14 March 2027", "Claims helpline 1800 209 4400"]))
    pages = extract_pages(path)
    assert [p.number for p in pages] == [1, 2]
    assert "14 March 2027" in pages[0].text and "1800 209 4400" in pages[1].text
    assert not any(p.ocr for p in pages)


def test_docx_includes_tables(tmp_path):
    d = docx.Document()
    d.add_paragraph("Warranty card for washing machine")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Serial Number"
    table.rows[0].cells[1].text = "VFL7012"
    d.save(tmp_path / "card.docx")
    [page] = extract_pages(tmp_path / "card.docx")
    assert "washing machine" in page.text and "Serial Number | VFL7012" in page.text


def test_text_and_markdown(tmp_path):
    (tmp_path / "bill.md").write_text("# Bill\nDue   date:\t5 Sep\n\n\n\nTotal", encoding="utf-8")
    [page] = extract_pages(tmp_path / "bill.md")
    assert page.text == "# Bill\nDue date: 5 Sep\n\nTotal"


def test_image_without_tesseract_explains_how_to_fix(tmp_path, monkeypatch):
    monkeypatch.setattr(extract, "ocr_available", lambda: False)
    monkeypatch.setattr(ocr, "ocr_available", lambda: False)
    Image.new("RGB", (20, 20), "white").save(tmp_path / "card.png")
    (tmp_path / "phone_photo.jpg").write_bytes(b"not decodable")  # the install hint must still win
    for name in ("card.png", "phone_photo.jpg"):
        with pytest.raises(OcrUnavailableError, match="brew install tesseract"):
            extract_pages(tmp_path / name)


def test_scanned_pdf_without_tesseract_explains_how_to_fix(tmp_path, monkeypatch):
    monkeypatch.setattr(extract, "ocr_available", lambda: False)
    monkeypatch.setattr(ocr, "ocr_available", lambda: False)
    (tmp_path / "scan.pdf").write_bytes(make_pdf([""]))
    with pytest.raises(OcrUnavailableError):
        extract_pages(tmp_path / "scan.pdf")


def test_unsupported_type(tmp_path):
    (tmp_path / "sheet.xlsx").write_bytes(b"x")
    with pytest.raises(UnsupportedFileError):
        extract_pages(tmp_path / "sheet.xlsx")


def test_clean_text_repairs_encoding():
    assert clean_text("monarchies â€“ a federal") == "monarchies – a federal"


@pytest.mark.parametrize(
    "filename,text,expected",
    [
        ("scan001.pdf", "Policy Number SGI/TW. Premium Rs 1,842. Sum insured. Nominee", "insurance"),
        ("IMG_2231.jpg", "Warranty card. Serial No. Date of purchase. Service centre", "warranty"),
        ("aug.pdf", "Amount due Rs 2,486. Due date 18 Sep. Units consumed 312", "bill"),
        ("deed.pdf", "Sale deed registered at the Sub-Registrar. Carpet area 1085", "property"),
        ("notes.txt", "Grocery list: milk, eggs", "other"),
    ],
)
def test_guess_category(filename, text, expected):
    assert guess_category(filename, text) == expected


def test_files_are_recognised_by_their_contents_not_their_names(tmp_path):
    """Phones and chat apps send files with no extension, a wrong one, or as HEIC photos."""
    from PIL import Image

    from backend.parser.extract import file_type, sniff

    pdf = tmp_path / "Aadhaar"  # WhatsApp: no extension
    pdf.write_bytes(b"%PDF-1.7\n%\xe2\xe3\xcf\xd3\n")
    assert sniff(pdf) == ".pdf" and file_type(pdf, "Aadhaar") == ".pdf"
    photo = tmp_path / "bill.pdf"  # a photo with the wrong name
    Image.new("RGB", (20, 20)).save(photo, "JPEG")
    assert file_type(photo, "bill.pdf") == ".jpg"
    heic = tmp_path / "IMG_2031.HEIC"
    import pillow_heif

    pillow_heif.register_heif_opener()
    Image.new("RGB", (64, 64), "white").save(heic, "HEIF")
    assert file_type(heic, heic.name) == ".heic"
    note = tmp_path / "note"
    note.write_text("Water bill due 15 July")
    assert file_type(note, "note") == ".txt"
    old_word = tmp_path / "letter.doc"
    old_word.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 100)
    with pytest.raises(UnsupportedFileError, match="save it as PDF"):
        file_type(old_word, "letter.doc")
