"""Highlighting the answer on the page it came from."""

from pathlib import Path

import docx
from fastapi.testclient import TestClient

from backend.api.app import create_app
from backend.chunker.chunker import RegexTokenizer
from backend.viewer.highlight import answer_facts, highlights_for
from backend.vector_store.store import DocumentStore
from tests.fakes import FakeEmbedder
from tests.pdf_helper import make_pdf

SAMPLES = Path(__file__).resolve().parent.parent / "samples" / "pdfs"
BILL = SAMPLES / "bill" / "Electricity Bill 2026 08.pdf"
ANSWER = "Your electricity bill for August is **Rs. 2,486** [1], due on 18 September 2026 [1]."


def test_answer_facts():
    assert answer_facts(ANSWER) == ["Rs. 2,486", "18 September 2026"]
    assert answer_facts("Policy SGI/TW/2026/448812 expires on 14 March 2027 [2].") == ["SGI/TW/2026/448812", "14 March 2027"]
    assert answer_facts("The carpet area is 1,085 sq ft.") == ["1,085 sq ft"]
    assert answer_facts("It renews in 2027, 5 days early.") == ["5 days"]  # a bare year alone is too vague
    assert answer_facts("Room 4 on floor 2.") == []  # so are tiny numbers
    assert answer_facts("I couldn't find that in your documents.") == []


def test_pdf_highlights_find_passage_and_facts_on_the_page():
    from backend.parser.extract import extract_pages

    page_text = extract_pages(BILL)[0].text
    h = highlights_for(BILL, 1, page_text, ANSWER)
    assert h.kind == "pdf" and h.pages == 1 and (h.width, h.height) == (612.0, 792.0)
    assert h.passage_found and h.passage
    assert h.facts_found == ["Rs. 2,486", "18 September 2026"] and not h.facts_missing
    amount = h.facts[0]
    x, y, w, hgt = amount
    assert 0 < x < 612 and 0 < y < 792 and w > 20 and hgt < 20  # one box around "Rs. 2,486" on a single line


def test_facts_not_on_the_page_are_reported():
    h = highlights_for(BILL, 1, None, "It is due on 5 Sept 2026 and costs Rs. 9,999.")
    assert h.facts == [] and h.facts_missing == ["5 Sept 2026", "Rs. 9,999"]


def test_scanned_pages_and_images_have_no_highlights(tmp_path):
    (tmp_path / "scan.pdf").write_bytes(make_pdf([""]))
    assert highlights_for(tmp_path / "scan.pdf", 1, "x", ANSWER).scanned
    assert highlights_for(tmp_path / "card.jpg", 1, "x", ANSWER).kind == "image"


def test_word_files_get_character_spans(tmp_path):
    d = docx.Document()
    d.add_paragraph("Water bill for April to June.")
    d.add_paragraph("Total amount due: Rs. 1,683. Due date: 15 July 2026.")
    d.save(tmp_path / "water.docx")
    h = highlights_for(tmp_path / "water.docx", 1, "Total amount due: Rs. 1,683.", "It is Rs. 1,683, due 15 July 2026.")
    assert h.kind == "text" and h.passage_found
    assert [h.text[s:e] for s, e in h.facts] == ["Rs. 1,683", "15 July 2026"]


def test_page_image_and_highlight_endpoints(tmp_path):
    client = TestClient(create_app(DocumentStore(tmp_path / "index"), FakeEmbedder(), RegexTokenizer()))
    doc = client.post("/api/documents", files=[("files", (BILL.name, BILL.read_bytes()))]).json()[0]["document"]
    chunk_id = client.get("/api/search", params={"q": "electricity bill"}).json()["results"][0]["chunk_id"]

    image = client.get(f"/api/documents/{doc['doc_id']}/pages/1/image")
    assert image.status_code == 200 and image.content.startswith(b"\x89PNG")
    assert image.headers["cache-control"] == "private, no-store"
    h = client.post(f"/api/documents/{doc['doc_id']}/pages/1/highlights",
                    json={"chunk_id": chunk_id, "answer": ANSWER}).json()
    assert h["passage_found"] and h["facts_found"] == ["Rs. 2,486", "18 September 2026"]

    assert client.get(f"/api/documents/{doc['doc_id']}/pages/2/image").status_code == 404
    assert client.get("/api/documents/nope/pages/1/image").status_code == 404
    txt = client.post("/api/documents", files=[("files", ("a.txt", b"Fridge warranty 10 years"))]).json()[0]["document"]
    assert client.get(f"/api/documents/{txt['doc_id']}/pages/1/image").status_code == 400
